import streamlit as st
from pathlib import Path

from utils.resume_rewriter import rewrite_resume
from utils.analytics import increase_resume_rewrite


def create_docx(resume_text: str, output_path: str):
    """Create a professional DOCX from rewritten resume text."""

    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    section = doc.sections[0]
    section.top_margin = Pt(40)
    section.bottom_margin = Pt(40)
    section.left_margin = Pt(50)
    section.right_margin = Pt(50)

    styles = doc.styles["Normal"]
    styles.font.name = "Arial"
    styles.font.size = Pt(10)

    lines = resume_text.splitlines()

    for line in lines:

        line = line.strip()

        if not line:
            continue

        # Heading detection
        if (
            line.isupper()
            or line.endswith(":")
            or line.lower() in {
                "summary",
                "professional summary",
                "experience",
                "education",
                "skills",
                "projects",
                "certifications",
                "achievements",
                "objective",
            }
        ):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(11)

        else:
            p = doc.add_paragraph(line)

            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(10)

    doc.save(output_path)


def show_resume_rewriter_page():

    st.header("✍ AI Resume Rewriter")

    st.write(
        "Rewrite your resume according to a specific Job Description."
    )

    resume = st.file_uploader(
        "Upload Resume",
        type=["pdf", "docx"],
        key="rewrite_resume"
    )

    job_description = st.text_area(
        "📋 Job Description",
        placeholder="Paste the complete job description here...",
        height=250
    )

    if not resume:
        return

    if not job_description.strip():
        st.info(
            "Paste the Job Description to create a targeted resume."
        )
        return

    upload_dir = Path("uploads")
    output_dir = Path("outputs")

    upload_dir.mkdir(exist_ok=True)
    output_dir.mkdir(exist_ok=True)

    resume_path = upload_dir / resume.name

    with open(resume_path, "wb") as file:
        file.write(resume.getbuffer())

    if st.button(
        "🚀 Rewrite Resume",
        use_container_width=True
    ):

        with st.spinner(
            "AI is tailoring your resume to the Job Description..."
        ):

            try:

                rewritten_resume = rewrite_resume(
                    str(resume_path),
                    job_description
                )

                increase_resume_rewrite()

                output_path = (
                    output_dir / "AI_Rewritten_Resume.docx"
                )

                create_docx(
                    rewritten_resume,
                    str(output_path)
                )

                st.session_state["rewritten_resume"] = (
                    rewritten_resume
                )

                st.session_state["rewritten_docx"] = (
                    str(output_path)
                )

                st.success(
                    "✅ Resume successfully rewritten!"
                )

            except Exception as error:

                st.error(
                    f"❌ Resume rewriting failed: {error}"
                )

    if "rewritten_resume" in st.session_state:

        st.subheader("📄 Rewritten Resume")

        st.text_area(
            "Preview",
            st.session_state["rewritten_resume"],
            height=500
        )

        docx_path = st.session_state["rewritten_docx"]

        with open(docx_path, "rb") as file:

            st.download_button(
                label="⬇ Download Rewritten Resume (.docx)",
                data=file,
                file_name="AI_Rewritten_Resume.docx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                use_container_width=True
            )