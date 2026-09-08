"""
DOCX Editor Utility
Extracts text from a resume and replaces selected paragraphs
while preserving the original DOCX template and paragraph formatting.
"""

import json
from pathlib import Path

from docx import Document

from config.logging_config import setup_logger

logger = setup_logger()


def extract_docx_text(file_path: str) -> str:
    """Extract resume text with paragraph IDs for AI rewriting."""

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Resume not found: {file_path}")

    if path.suffix.lower() != ".docx":
        raise ValueError("Only DOCX files are supported.")

    document = Document(file_path)
    lines = []

    paragraph_id = 1

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            lines.append(f"[P{paragraph_id}] {text}")

        paragraph_id += 1

    for table_index, table in enumerate(document.tables, start=1):

        for row_index, row in enumerate(table.rows, start=1):

            for col_index, cell in enumerate(row.cells, start=1):

                text = cell.text.strip()

                if text:
                    lines.append(
                        f"[T{table_index}_{row_index}_{col_index}] {text}"
                    )

    result = "\n".join(lines)

    if not result.strip():
        raise ValueError("No readable text found in DOCX.")

    logger.info("DOCX text extracted successfully.")

    return result


def _copy_run_format(source_run, target_run):
    """Copy basic formatting from an existing run."""

    target_run.bold = source_run.bold
    target_run.italic = source_run.italic
    target_run.underline = source_run.underline

    if source_run.font.name:
        target_run.font.name = source_run.font.name

    if source_run.font.size:
        target_run.font.size = source_run.font.size

    if source_run.font.color and source_run.font.color.type:
        try:
            target_run.font.color.rgb = source_run.font.color.rgb
        except Exception:
            pass


def _replace_paragraph_text(paragraph, new_text: str):
    """Replace paragraph text while retaining basic original formatting."""

    if not paragraph.runs:
        paragraph.add_run(new_text)
        return

    first_run = paragraph.runs[0]

    for run in paragraph.runs:
        run.text = ""

    first_run.text = new_text

    for run in paragraph.runs[1:]:
        run.text = ""

    _copy_run_format(first_run, first_run)


def _replace_table_cell(cell, new_text: str):
    """Replace text inside a table cell."""

    if cell.paragraphs:
        _replace_paragraph_text(
            cell.paragraphs[0],
            new_text
        )


def apply_rewrites(
    input_path: str,
    output_path: str,
    replacements: dict
) -> str:
    """
    Apply AI-generated paragraph replacements to the original DOCX.

    Only paragraphs returned by the AI are changed.
    The original document structure remains intact.
    """

    if not Path(input_path).exists():
        raise FileNotFoundError(
            f"Input DOCX not found: {input_path}"
        )

    document = Document(input_path)

    # ---------------------------------------------
    # Normal paragraphs
    # ---------------------------------------------

    for index, paragraph in enumerate(
        document.paragraphs,
        start=1
    ):

        key = f"P{index}"

        if key in replacements:

            new_text = str(
                replacements[key]
            ).strip()

            if new_text:
                _replace_paragraph_text(
                    paragraph,
                    new_text
                )

    # ---------------------------------------------
    # Tables
    # ---------------------------------------------

    for table_index, table in enumerate(
        document.tables,
        start=1
    ):

        for row_index, row in enumerate(
            table.rows,
            start=1
        ):

            for col_index, cell in enumerate(
                row.cells,
                start=1
            ):

                key = (
                    f"T{table_index}_"
                    f"{row_index}_"
                    f"{col_index}"
                )

                if key in replacements:

                    new_text = str(
                        replacements[key]
                    ).strip()

                    if new_text:
                        _replace_table_cell(
                            cell,
                            new_text
                        )

    Path(output_path).parent.mkdir(
        parents=True,
        exist_ok=True
    )

    document.save(output_path)

    logger.info(
        f"Optimized DOCX saved: {output_path}"
    )

    return output_path


def parse_rewrite_response(response: str) -> dict:
    """Extract the replacement dictionary from AI response."""

    response = response.strip()

    # Remove markdown code fences
    if response.startswith("```"):
        lines = response.splitlines()

        if lines and lines[0].startswith("```"):
            lines = lines[1:]

        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]

        response = "\n".join(lines).strip()

    try:
        data = json.loads(response)

    except json.JSONDecodeError as error:

        logger.exception(error)

        raise ValueError(
            "AI returned invalid JSON."
        ) from error

    replacements = data.get(
        "replacements"
    )

    if not isinstance(
        replacements,
        dict
    ):
        raise ValueError(
            "AI response does not contain "
            "'replacements'."
        )

    return replacements